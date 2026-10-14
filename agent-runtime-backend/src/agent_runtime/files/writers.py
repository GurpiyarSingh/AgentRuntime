"""Creating a file of (almost) any type from text a model wrote.

The extension picks the format:

- `.docx` and `.pdf` are rendered from light Markdown: `#` headings,
  `-`/`1.` lists, `|` tables and `**bold**`.
- `.xlsx` takes CSV, a JSON list of rows (lists or objects), or a JSON
  object mapping sheet names to rows. A cell starting with `=` is a formula.
- `.pptx` takes slides separated by a `---` line: the first line is the
  title, `- ` lines are bullets (indent two spaces per level), and a line
  starting `Notes:` becomes the speaker notes.
- Any other extension is written as UTF-8 text (`.json` is validated).
- `encoding="base64"` writes the decoded bytes verbatim, for images,
  archives or any other binary format.

Every write goes through a temp file, so a failure never leaves a
half-written document behind.
"""

from __future__ import annotations

import base64
import binascii
import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

import docx
import openpyxl
from fpdf import FPDF, XPos, YPos
from openpyxl.styles import Font
from pptx import Presentation

from agent_runtime.files.images import RASTER_EXTENSIONS
from agent_runtime.files.workspace import atomic_target

Encoding = Literal["text", "base64"]

# Extensions whose content is never text, so writing a model's prose into
# them would only produce a corrupt file.
BINARY_ONLY_EXTENSIONS = RASTER_EXTENSIONS | {
    ".zip", ".gz", ".tgz", ".bz2", ".xz", ".7z", ".rar", ".tar", ".exe", ".dll",
    ".so", ".dylib", ".bin", ".mp3", ".mp4", ".wav", ".ogg", ".flac", ".avi",
    ".mov", ".mkv", ".webm", ".woff", ".woff2", ".ttf", ".otf", ".doc", ".xls", ".ppt",
}  # fmt: skip


class WriteError(Exception):
    """The content can't be turned into the requested format."""


# --- light Markdown --------------------------------------------------------------


@dataclass(slots=True)
class Block:
    """One structural element of a document: a heading, paragraph, list item or table."""

    kind: Literal["heading", "paragraph", "bullet", "numbered", "table"]
    text: str = ""
    level: int = 0
    rows: list[list[str]] = field(default_factory=list)


_HEADING = re.compile(r"(#{1,6})\s+(.*)")
_BULLET = re.compile(r"[-*+]\s+(.*)")
_NUMBERED = re.compile(r"\d+[.)]\s+(.*)")
_TABLE_RULE = re.compile(r":?-{2,}:?")


def parse_markdown(content: str) -> list[Block]:
    """Split text into the handful of block types the document writers render."""
    blocks: list[Block] = []
    paragraph: list[str] = []
    table: list[list[str]] = []

    def flush() -> None:
        if paragraph:
            blocks.append(Block("paragraph", "\n".join(paragraph)))
            paragraph.clear()
        if table:
            width = max(len(row) for row in table)
            blocks.append(Block("table", rows=[row + [""] * (width - len(row)) for row in table]))
            table.clear()

    for raw in content.splitlines():
        line = raw.strip()
        if len(line) > 1 and line.startswith("|") and line.endswith("|"):
            if paragraph:
                flush()
            cells = [cell.strip() for cell in line[1:-1].split("|")]
            if not all(_TABLE_RULE.fullmatch(cell) for cell in cells):
                table.append(cells)
            continue
        if table:
            flush()
        if not line:
            flush()
        elif match := _HEADING.fullmatch(line):
            flush()
            blocks.append(Block("heading", match.group(2).strip(), level=len(match.group(1))))
        elif match := _BULLET.fullmatch(line):
            flush()
            blocks.append(Block("bullet", match.group(1)))
        elif match := _NUMBERED.fullmatch(line):
            flush()
            blocks.append(Block("numbered", match.group(1)))
        else:
            paragraph.append(line)
    flush()
    return blocks


def _bold_runs(text: str) -> list[tuple[str, bool]]:
    """`a **b** c` -> [("a ", False), ("b", True), (" c", False)]."""
    return [(part, index % 2 == 1) for index, part in enumerate(text.split("**")) if part]


# --- DOCX ------------------------------------------------------------------------


def _write_docx(content: str, target: Path) -> None:
    document = docx.Document()
    for block in parse_markdown(content):
        if block.kind == "heading":
            document.add_heading(block.text, level=min(block.level, 9))
        elif block.kind == "table":
            table = document.add_table(rows=0, cols=len(block.rows[0]), style="Table Grid")
            for row_index, row in enumerate(block.rows):
                cells = table.add_row().cells
                for cell, value in zip(cells, row, strict=True):
                    cell.text = value
                    if row_index == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
        else:
            style = {"bullet": "List Bullet", "numbered": "List Number"}.get(block.kind)
            paragraph = document.add_paragraph(style=style)
            for text, bold in _bold_runs(block.text):
                paragraph.add_run(text).bold = bold
    document.save(str(target))


# --- PDF -------------------------------------------------------------------------

# Core PDF fonts only cover Latin-1. A system TrueType font, when one is
# found, lets the PDF carry any script the font has glyphs for.
_FONT_CANDIDATES = [
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("C:/Windows/Fonts/segoeui.ttf", "C:/Windows/Fonts/segoeuib.ttf"),
    (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    ),
    ("/usr/share/fonts/TTF/DejaVuSans.ttf", "/usr/share/fonts/TTF/DejaVuSans-Bold.ttf"),
    (
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    ),
]
_ASCII_FALLBACKS = str.maketrans(
    {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-",
     "\u2014": "-", "\u2026": "...", "\u2022": "-", "\u00a0": " "}
)  # fmt: skip


def _pdf_font(pdf: FPDF) -> tuple[str, bool]:
    """Register the best available font; returns (family, is_unicode)."""
    for regular, bold in _FONT_CANDIDATES:
        if Path(regular).is_file():
            pdf.add_font("Body", "", regular)
            pdf.add_font("Body", "B", bold if Path(bold).is_file() else regular)
            return "Body", True
    return "Helvetica", False


def _write_pdf(content: str, target: Path) -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    family, unicode_font = _pdf_font(pdf)

    def clean(text: str) -> str:
        if unicode_font:
            return text
        return text.translate(_ASCII_FALLBACKS).encode("latin-1", "replace").decode("latin-1")

    def write(text: str, size: float, height: float, bold: bool = False, indent: float = 0) -> None:
        pdf.set_font(family, "B" if bold else "", size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(
            0, height, clean(text), markdown=not bold, new_x=XPos.LMARGIN, new_y=YPos.NEXT
        )

    number = 0
    for block in parse_markdown(content):
        number = number + 1 if block.kind == "numbered" else 0
        if block.kind == "heading":
            pdf.ln(2)
            write(block.text, size=max(20 - 3 * (block.level - 1), 12), height=9, bold=True)
            pdf.ln(1)
        elif block.kind == "paragraph":
            write(block.text, size=11, height=6)
            pdf.ln(3)
        elif block.kind == "bullet":
            write(("\u2022 " if unicode_font else "- ") + block.text, 11, 6, indent=5)
        elif block.kind == "numbered":
            write(f"{number}. {block.text}", 11, 6, indent=5)
        elif block.kind == "table":
            pdf.set_font(family, "", 10)
            with pdf.table() as table:
                for row in block.rows:
                    cells = table.row()
                    for value in row:
                        cells.cell(clean(value))
            pdf.ln(3)
    pdf.output(str(target))


# --- XLSX ------------------------------------------------------------------------

_SHEET_NAME_INVALID = re.compile(r"[\[\]:*?/\\]")


def _coerce(value: str) -> Any:
    """CSV gives strings; keep numbers numeric so the sheet can compute with them."""
    text = value.strip()
    for cast in (int, float):
        try:
            return cast(text)
        except ValueError:
            continue
    return value


def _rows(data: Any) -> list[list[Any]]:
    if not isinstance(data, list):
        raise WriteError("Spreadsheet rows must be a JSON list.")
    if data and all(isinstance(row, dict) for row in data):
        header = list(dict.fromkeys(key for row in data for key in row))
        return [header, *([row.get(key) for key in header] for row in data)]
    return [row if isinstance(row, list) else [row] for row in data]


def _sheets(content: str) -> dict[str, list[list[Any]]]:
    stripped = content.strip()
    if stripped[:1] in ("{", "["):
        try:
            data = json.loads(stripped)
        except json.JSONDecodeError:
            data = None
        if isinstance(data, dict):
            return {str(name): _rows(rows) for name, rows in data.items()}
        if data is not None:
            return {"Sheet1": _rows(data)}
    rows = csv.reader(io.StringIO(content))
    return {"Sheet1": [[_coerce(cell) for cell in row] for row in rows]}


def _cell(value: Any) -> Any:
    if isinstance(value, dict | list):
        return json.dumps(value)
    return value


def _write_xlsx(content: str, target: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]
    for name, rows in _sheets(content).items():
        sheet = workbook.create_sheet(_SHEET_NAME_INVALID.sub("_", name)[:31] or "Sheet")
        for row in rows:
            sheet.append([_cell(value) for value in row])
        for cell in sheet[1] if rows else ():
            cell.font = Font(bold=True)
    if not workbook.worksheets:
        workbook.create_sheet("Sheet1")
    workbook.save(target)


# --- PPTX ------------------------------------------------------------------------

_SLIDE_BREAK = re.compile(r"^\s*---+\s*$", re.MULTILINE)
_TITLE_AND_CONTENT, _TITLE_ONLY = 1, 5


def _is_notes(line: str) -> bool:
    return line.strip().startswith("Notes:")


def _write_pptx(content: str, target: Path) -> None:
    deck = Presentation()
    for chunk in _SLIDE_BREAK.split(content):
        lines = [line.rstrip() for line in chunk.splitlines() if line.strip()]
        if not lines:
            continue
        title = lines[0].strip().lstrip("#").strip()
        notes = [line.strip()[6:].strip() for line in lines[1:] if _is_notes(line)]
        body = [line for line in lines[1:] if not _is_notes(line)]

        layout = deck.slide_layouts[_TITLE_AND_CONTENT if body else _TITLE_ONLY]
        slide = deck.slides.add_slide(layout)
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        if body:
            frame = slide.placeholders[1].text_frame
            for index, line in enumerate(body):
                indent = (len(line) - len(line.lstrip())) // 2
                text = line.strip()
                if match := _BULLET.fullmatch(text):
                    text = match.group(1)
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = text
                paragraph.level = min(indent, 4)
        if notes:
            slide.notes_slide.notes_text_frame.text = "\n".join(notes)  # type: ignore[union-attr]
    deck.save(str(target))


# --- dispatch ------------------------------------------------------------------------


_RENDERERS = {
    ".docx": ("docx", _write_docx),
    ".pdf": ("pdf", _write_pdf),
    ".xlsx": ("xlsx", _write_xlsx),
    ".pptx": ("pptx", _write_pptx),
}


def _decode_base64(content: str) -> bytes:
    payload = content.strip()
    if payload.startswith("data:") and "," in payload:
        payload = payload.split(",", 1)[1]
    try:
        return base64.b64decode("".join(payload.split()), validate=True)
    except (binascii.Error, ValueError) as exc:
        raise WriteError(f"Content is not valid base64: {exc}") from exc


def write_file(path: Path, content: str, encoding: Encoding = "text") -> str:
    """Create `path` from `content`; returns the format it was written as."""
    suffix = path.suffix.lower()
    with atomic_target(path) as target:
        if encoding == "base64":
            target.write_bytes(_decode_base64(content))
            return "binary"
        if suffix in BINARY_ONLY_EXTENSIONS:
            raise WriteError(
                f"{suffix} is a binary format. Pass its bytes with encoding='base64', "
                "or choose a text-based format."
            )
        if suffix in _RENDERERS:
            kind, render = _RENDERERS[suffix]
            render(content, target)
            return kind
        if suffix == ".json":
            try:
                json.loads(content)
            except json.JSONDecodeError as exc:
                raise WriteError(f"Content is not valid JSON: {exc}") from exc
        target.write_text(content, encoding="utf-8", newline="")
        return "text"


__all__ = ["BINARY_ONLY_EXTENSIONS", "Block", "WriteError", "parse_markdown", "write_file"]
