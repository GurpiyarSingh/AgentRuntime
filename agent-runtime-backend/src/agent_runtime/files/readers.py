"""Turning a file of (almost) any type into text a model can read.

Each format gets the most faithful plain-text rendering that is cheap to
produce: documents keep their headings, lists and tables; spreadsheets
become CSV per sheet; slide decks keep slide order and speaker notes;
archives are listed rather than unpacked. Anything with no textual form
is described (type, size, leading bytes) instead of dumped as noise.

Images are deliberately not handled here — they are *looked at*, not
read, through `agent_runtime.files.images` and the `read_image` tool.
"""

from __future__ import annotations

import codecs
import csv
import io
import mimetypes
import tarfile
import zipfile
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import docx
import openpyxl
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

from agent_runtime.files.images import RASTER_EXTENSIONS

ARCHIVE_LISTING_LIMIT = 500
_TAR_SUFFIXES = (".tar", ".tar.gz", ".tgz", ".tar.bz2", ".tbz2", ".tar.xz", ".txz")
_LEGACY_OFFICE = {
    ".doc": "Word 97-2003 (.doc)",
    ".xls": "Excel 97-2003 (.xls)",
    ".ppt": "PowerPoint 97-2003 (.ppt)",
}


class DocumentError(Exception):
    """The file is corrupt, encrypted, or otherwise unreadable as its type."""


@dataclass(frozen=True, slots=True)
class Document:
    """A file's text, and what kind of file it was read as."""

    kind: str
    text: str


# --- plain text ----------------------------------------------------------------


def decode_text(data: bytes) -> str | None:
    """Decode bytes as text, or `None` if they look binary.

    BOMs are honoured; otherwise UTF-8 is tried first and Windows-1252 is
    the fallback for legacy files. NUL bytes or a high share of control
    characters mean binary, so an executable is never "read" as gibberish.
    """
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8) :].decode("utf-8", errors="replace")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace")
    sample = data[:8192]
    if b"\x00" in sample:
        return None
    control = sum(1 for byte in sample if byte < 32 and byte not in b"\t\n\r\f\b")
    if sample and control / len(sample) > 0.1:
        return None
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


# --- documents -----------------------------------------------------------------


def _read_pdf(path: Path) -> str:
    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentError("The PDF is password-protected.")
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        pages.append(
            f"--- Page {number} ---\n"
            + (text or "[no extractable text on this page; it may be a scanned image]")
        )
    return "\n\n".join(pages) if pages else "[the PDF has no pages]"


def _table_rows(rows: Iterable[Iterable[str]]) -> list[str]:
    return ["| " + " | ".join(cell.replace("\n", " ") for cell in row) + " |" for row in rows]


def _read_docx(path: Path) -> str:
    document = docx.Document(str(path))
    lines: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            lines.extend(_table_rows([cell.text for cell in row.cells] for row in block.rows))
            lines.append("")
            continue
        if not isinstance(block, Paragraph):  # pragma: no cover - python-docx yields only these
            continue
        text = block.text
        style = block.style.name if block.style is not None else ""
        if style == "Title":
            lines.append(f"# {text}")
        elif style.startswith("Heading") and style[-1:].isdigit():
            lines.append("#" * min(int(style[-1]), 6) + f" {text}")
        elif "List" in style and text:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    return "\n".join(lines).strip()


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def _read_xlsx(path: Path) -> str:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sections = []
        for sheet in workbook.worksheets:
            buffer = io.StringIO()
            writer = csv.writer(buffer, lineterminator="\n")
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell_text(value) for value in row]
                if any(cells):
                    writer.writerow(cells)
            body = buffer.getvalue().strip() or "[empty sheet]"
            sections.append(f"## Sheet: {sheet.title}\n{body}")
        return "\n\n".join(sections)
    finally:
        workbook.close()


def _shape_text(shape: Any) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [line for child in shape.shapes for line in _shape_text(child)]
    if getattr(shape, "has_table", False) and shape.has_table:
        return _table_rows([cell.text for cell in row.cells] for row in shape.table.rows)
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text.strip()
        return [text] if text else []
    return []


def _read_pptx(path: Path) -> str:
    deck = Presentation(str(path))
    slides = []
    for number, slide in enumerate(deck.slides, start=1):
        lines = [line for shape in slide.shapes for line in _shape_text(shape)]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                lines.append(f"Notes: {notes.text.strip()}")
        slides.append(f"--- Slide {number} ---\n" + ("\n".join(lines) or "[no text]"))
    return "\n\n".join(slides) if slides else "[the presentation has no slides]"


# --- archives and binaries ---------------------------------------------------------


def _listing(entries: list[tuple[str, int]]) -> str:
    shown = entries[:ARCHIVE_LISTING_LIMIT]
    lines = [f"{name}  ({size} bytes)" for name, size in shown]
    if len(entries) > len(shown):
        lines.append(f"... and {len(entries) - len(shown)} more entries")
    return f"{len(entries)} entries:\n" + "\n".join(lines)


def _read_zip(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return _listing(
            [(info.filename, info.file_size) for info in archive.infolist() if not info.is_dir()]
        )


def _read_tar(path: Path) -> str:
    with tarfile.open(path) as archive:
        return _listing([(m.name, m.size) for m in archive.getmembers() if m.isfile()])


def _describe_binary(path: Path, data: bytes) -> str:
    """Summarise a file with no text form: its likely type, size and leading bytes."""
    mime = mimetypes.guess_type(path.name)[0] or "unknown type"
    head = data[:64].hex(" ")
    return (
        f"[binary file: {mime}, {len(data)} bytes — it has no text form to show]\n"
        f"First {min(len(data), 64)} bytes (hex): {head}"
    )


# --- dispatch ------------------------------------------------------------------------


_EXTRACTORS: dict[str, tuple[str, Callable[[Path], str]]] = {
    ".pdf": ("pdf", _read_pdf),
    ".docx": ("docx", _read_docx),
    ".docm": ("docx", _read_docx),
    ".xlsx": ("xlsx", _read_xlsx),
    ".xlsm": ("xlsx", _read_xlsx),
    ".pptx": ("pptx", _read_pptx),
    ".pptm": ("pptx", _read_pptx),
}


def read_document(path: Path) -> Document:
    """Read any file as text. Raises `DocumentError` if its format is broken."""
    suffix = path.suffix.lower()
    name = path.name.lower()

    if suffix in RASTER_EXTENSIONS:
        raise DocumentError(f"{path.name} is an image. Use the read_image tool to look at it.")
    if suffix in _LEGACY_OFFICE:
        raise DocumentError(
            f"{_LEGACY_OFFICE[suffix]} is a legacy binary format that can't be read. "
            f"Ask for it to be saved as {suffix}x."
        )

    extractor = _EXTRACTORS.get(suffix)
    data = b""
    if extractor is None:
        data = path.read_bytes()
        if data.startswith(b"%PDF-"):
            extractor = _EXTRACTORS[".pdf"]

    try:
        if extractor is not None:
            kind, extract = extractor
            return Document(kind=kind, text=extract(path))
        if name.endswith(_TAR_SUFFIXES):
            return Document(kind="tar archive", text=_read_tar(path))
        text = decode_text(data)
        if text is not None:
            return Document(kind="text", text=text)
        if zipfile.is_zipfile(path):
            return Document(kind="zip archive", text=_read_zip(path))
        if tarfile.is_tarfile(path):
            return Document(kind="tar archive", text=_read_tar(path))
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(
            f"Could not read {path.name} as {suffix or 'its type'}: {type(exc).__name__}: {exc}"
        ) from exc

    return Document(kind="binary", text=_describe_binary(path, data))


__all__ = ["Document", "DocumentError", "decode_text", "read_document"]
