"""Reading and writing every supported format, mostly as round trips.

Writing a document with `write_file` and reading it back with
`read_document` checks both directions at once against real libraries,
with no network and nothing outside `tmp_path`.
"""

from __future__ import annotations

import base64
import io
import json
import tarfile
import zipfile
from pathlib import Path

import openpyxl
import pytest
from PIL import Image

from agent_runtime.files.readers import DocumentError, decode_text, read_document
from agent_runtime.files.writers import WriteError, parse_markdown, write_file

# --- text decoding ---------------------------------------------------------------


def test_decode_text_handles_utf8_boms_and_legacy_encodings() -> None:
    assert decode_text("héllo".encode()) == "héllo"
    assert decode_text(b"\xef\xbb\xbfbom") == "bom"
    assert decode_text("utf16".encode("utf-16")) == "utf16"
    assert decode_text("café".encode("cp1252")) == "café"


def test_decode_text_refuses_binary() -> None:
    assert decode_text(b"\x00\x01\x02binary") is None
    assert decode_text(bytes(range(1, 32)) * 50) is None


# --- plain files -----------------------------------------------------------------


def test_text_round_trip_preserves_content_exactly(tmp_path: Path) -> None:
    path = tmp_path / "script.py"
    source = "def f():\r\n    return 'ünïcode'\r\n"
    assert write_file(path, source) == "text"
    assert path.read_bytes() == source.encode()
    assert read_document(path).text == source


def test_json_is_validated(tmp_path: Path) -> None:
    write_file(tmp_path / "ok.json", '{"a": 1}')
    with pytest.raises(WriteError, match="not valid JSON"):
        write_file(tmp_path / "bad.json", "{a: 1}")
    assert not (tmp_path / "bad.json").exists()


def test_binary_formats_need_base64(tmp_path: Path) -> None:
    with pytest.raises(WriteError, match="base64"):
        write_file(tmp_path / "pic.png", "not an image")


def test_base64_writes_exact_bytes(tmp_path: Path) -> None:
    payload = bytes(range(256))
    encoded = base64.b64encode(payload).decode()
    assert write_file(tmp_path / "blob.bin", encoded, encoding="base64") == "binary"
    assert (tmp_path / "blob.bin").read_bytes() == payload
    # A data: URL prefix is tolerated, as models often produce one.
    write_file(tmp_path / "blob2.bin", f"data:application/octet-stream;base64,{encoded}", "base64")
    assert (tmp_path / "blob2.bin").read_bytes() == payload


def test_invalid_base64_is_rejected(tmp_path: Path) -> None:
    with pytest.raises(WriteError, match="base64"):
        write_file(tmp_path / "x.bin", "!!! not base64 !!!", encoding="base64")


def test_unknown_binary_is_described_not_dumped(tmp_path: Path) -> None:
    path = tmp_path / "data.bin"
    path.write_bytes(b"\x00\x01\x02\x03" * 100)
    document = read_document(path)
    assert document.kind == "binary"
    assert "400 bytes" in document.text
    assert "00 01 02 03" in document.text


def test_images_are_redirected_to_read_image(tmp_path: Path) -> None:
    path = tmp_path / "photo.jpg"
    Image.new("RGB", (4, 4)).save(path)
    with pytest.raises(DocumentError, match="read_image"):
        read_document(path)


def test_legacy_office_formats_get_a_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "old.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(DocumentError, match="legacy"):
        read_document(path)


def test_corrupt_documents_raise_document_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_text("this is not a zip")
    with pytest.raises(DocumentError, match=r"broken.docx"):
        read_document(path)


# --- archives --------------------------------------------------------------------


def test_zip_archives_are_listed(tmp_path: Path) -> None:
    path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("a.txt", "hello")
        archive.writestr("dir/b.txt", "world!")
    document = read_document(path)
    assert document.kind == "zip archive"
    assert "a.txt  (5 bytes)" in document.text
    assert "dir/b.txt  (6 bytes)" in document.text


def test_tar_archives_are_listed(tmp_path: Path) -> None:
    path = tmp_path / "bundle.tar.gz"
    with tarfile.open(path, "w:gz") as archive:
        data = b"content"
        info = tarfile.TarInfo("inner.txt")
        info.size = len(data)
        archive.addfile(info, io.BytesIO(data))
    document = read_document(path)
    assert document.kind == "tar archive"
    assert "inner.txt" in document.text


# --- markdown parsing ------------------------------------------------------------


def test_parse_markdown_recognises_every_block_type() -> None:
    blocks = parse_markdown(
        "# Title\n\nFirst line\nsecond line\n\n- bullet\n1. numbered\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n\nAfter"
    )
    assert [b.kind for b in blocks] == [
        "heading", "paragraph", "bullet", "numbered", "table", "paragraph",
    ]  # fmt: skip
    assert blocks[0].level == 1
    assert blocks[1].text == "First line\nsecond line"
    assert blocks[4].rows == [["A", "B"], ["1", "2"]]


# --- office documents and PDF ---------------------------------------------------------

MARKDOWN = """# Quarterly Report

Revenue grew **strongly** this quarter.

## Highlights

- New customers: 42
- Churn down

1. Hire
2. Ship

| Region | Sales |
|--------|-------|
| North  | 100   |
| South  | 80    |
"""


def test_docx_round_trip_keeps_structure(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    assert write_file(path, MARKDOWN) == "docx"

    document = read_document(path)
    assert document.kind == "docx"
    text = document.text
    assert "# Quarterly Report" in text
    assert "## Highlights" in text
    assert "Revenue grew strongly this quarter." in text
    assert "- New customers: 42" in text
    assert "| Region | Sales |" in text
    assert "| North | 100 |" in text


def test_pdf_round_trip(tmp_path: Path) -> None:
    path = tmp_path / "report.pdf"
    assert write_file(path, MARKDOWN + "\nCurly quotes: “hi” — done") == "pdf"
    assert path.read_bytes().startswith(b"%PDF-")

    document = read_document(path)
    assert document.kind == "pdf"
    assert "--- Page 1 ---" in document.text
    assert "Quarterly Report" in document.text
    assert "North" in document.text


def test_pdf_detected_by_content_without_extension(tmp_path: Path) -> None:
    write_file(tmp_path / "report.pdf", "Hello PDF")
    renamed = tmp_path / "report_no_ext"
    (tmp_path / "report.pdf").rename(renamed)
    assert read_document(renamed).kind == "pdf"


def test_xlsx_from_csv_keeps_numbers_numeric(tmp_path: Path) -> None:
    path = tmp_path / "sales.xlsx"
    assert write_file(path, "region,sales\nNorth,100\nSouth,80.5\nTotal,=SUM(B2:B3)") == "xlsx"

    sheet = openpyxl.load_workbook(path).active
    assert sheet is not None
    assert sheet["B2"].value == 100
    assert sheet["B3"].value == 80.5
    assert sheet["B4"].value == "=SUM(B2:B3)"
    assert sheet["A1"].font.bold

    text = read_document(path).text
    assert "## Sheet: Sheet1" in text
    assert "North,100" in text


def test_xlsx_from_json_with_multiple_sheets(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    content = json.dumps(
        {
            "People": [{"name": "Ada", "age": 36}, {"name": "Alan", "age": 41}],
            "Raw/Data": [[1, 2], [3, 4]],
        }
    )
    write_file(path, content)

    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ["People", "Raw_Data"]
    people = workbook["People"]
    assert [c.value for c in people[1]] == ["name", "age"]
    assert [c.value for c in people[3]] == ["Alan", 41]


def test_pptx_round_trip_with_bullets_and_notes(tmp_path: Path) -> None:
    path = tmp_path / "deck.pptx"
    content = (
        "# Welcome\nNotes: say hello\n---\nAgenda\n- Intro\n  - Details\n- Wrap up\n---\nThanks!"
    )
    assert write_file(path, content) == "pptx"

    text = read_document(path).text
    assert "--- Slide 1 ---\nWelcome" in text
    assert "Notes: say hello" in text
    for expected in ("Agenda", "Intro", "Details"):
        assert expected in text
    assert "--- Slide 3 ---\nThanks!" in text


def test_failed_render_leaves_no_file(tmp_path: Path) -> None:
    with pytest.raises(WriteError):
        write_file(tmp_path / "bad.xlsx", '{"Sheet": "not a list"}')
    assert list(tmp_path.iterdir()) == []
